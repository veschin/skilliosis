#!/usr/bin/env python3
"""Check that every explicitly created numeric variable participates in at least one formula.
Literals (text-valued cells) are exempt by design. Numeric variables that are passive by
design (the owner applies them manually) are exempted explicitly via --allow COORD.
Usage: check_usage.py WORKBOOK.xlsx [--allow B21]
Exit 0 = every numeric cell is used; exit 1 = unused cells listed with their labels."""
import re
import sys

import openpyxl

REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
RANGE = re.compile(r"(\$?[A-Z]{1,3}\$?\d+):(\$?[A-Z]{1,3}\$?\d+)")


def col_num(col):
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch) - 64)
    return n


def col_letters(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def refs_from_formula(f):
    out = set()
    for m in RANGE.finditer(f):
        c1 = re.match(r"([A-Z]+)(\d+)", m.group(1).replace("$", ""))
        c2 = re.match(r"([A-Z]+)(\d+)", m.group(2).replace("$", ""))
        for col in range(col_num(c1.group(1)), col_num(c2.group(1)) + 1):
            for row in range(int(c1.group(2)), int(c2.group(2)) + 1):
                out.add(f"{col_letters(col)}{row}")
    for m in REF.finditer(f):
        out.add(f"{m.group(1)}{m.group(2)}")
    return out


def main():
    args = sys.argv[1:]
    allow = set()
    while "--allow" in args:
        i = args.index("--allow")
        allow.add(args[i + 1].replace("$", ""))
        del args[i:i + 2]
    if not args:
        print(__doc__)
        sys.exit(2)
    ws = openpyxl.load_workbook(args[0]).active
    used, formulas = set(), 0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                formulas += 1
                used |= refs_from_formula(c.value)
    unused = []
    for row in ws.iter_rows():
        for c in row:
            if c.value is None or isinstance(c.value, str):
                continue  # text cells are literals/labels: exempt
            if c.coordinate in used or c.coordinate in allow:
                continue
            label = ws.cell(c.row, 1).value
            if not isinstance(label, str) or label.startswith("="):
                label = ""
            unused.append((c.coordinate, label, c.value))
    if unused:
        print("UNUSED NUMERIC VARIABLE CELLS (text literals exempt; --allow COORD for passive-by-design):")
        for coord, label, v in unused:
            print(f"  {coord}: {v!r}  ({label})")
        sys.exit(1)
    print(f"OK: all numeric cells participate in formulas ({formulas} formulas scanned); text literals exempt")


main()
