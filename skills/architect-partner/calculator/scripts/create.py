#!/usr/bin/env python3
"""Create the architectural calculator skeleton with the user's palette (NEW workbooks only;
never run on an existing book - it would touch widths). Usage: create.py OUTPUT.xlsx"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Color, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "calculator.xlsx"
TB = 1000000000000
MN = 1000000
KB = 1000

FILL_SECTION = PatternFill("solid", fgColor=Color(theme=1, tint=0.35))
FILL_SUM = PatternFill("solid", fgColor=Color(theme=6, tint=0.8))
FILL_LABEL = PatternFill("solid", fgColor=Color(theme=3, tint=0.8))
FONT_WHITE = Font(color=Color(theme=0))
FONT_TEXT = Font(color=Color(theme=1))
FONT_LITERAL = Font(color=Color(theme=1), italic=True)

wb = Workbook()
ws = wb.active
ws.title = "Калькулятор данных"

# (coordinate, value, kind) - kind in section|sum|label|value|literal|headercell
S = []
S += [("A1", "глобальные переменные", "section")]
S += [("A2", "индексы сжатия паркета", "section"), ("B2", "процент сжатия", "section")]
for i, (lvl, pct) in enumerate([("-1", 47.63), ("-3", 48.36), ("-9", 49.74), ("-19", 51.66)]):
    S += [(f"A{3+i}", f"zstd {lvl}, % сжатия", "label"), (f"B{3+i}", pct, "value")]
S += [("A7", "выбранный индекс сжатия паркета", "label"), ("B7", "zstd -3", "literal")]
S += [("A8", "процент сжатия при укладке в паркет, %", "label"), ("B8", 85.56, "value")]
S += [("A9", "множитель репликации данных в субд, раз", "label"), ("B9", 2, "value")]
S += [("A10", "индекс сжатия кафки, %", "label"), ("B10", 0, "value")]
S += [("A11", "процент сжатия базовый у субд, %", "label"), ("B11", 30, "value")]
S += [("A13", "переменные под источник", "section"), ("B13", "название источника", "section")]
S += [("B14", "Интеграция 1", "literal")]
S += [("A15", "записей в сутки, миллионы", "label"), ("B15", 1000, "value")]
S += [("A16", "вес записи, кб", "label"), ("B16", None, "value")]
S += [("A17", "множитель избыточности бизнес данных в субд, раз", "label"), ("B17", 10, "value")]
S += [("A18", "процент записи фх который нужен субд, %", "label"), ("B18", 40, "value")]
S += [("A20", "вес записи в фх, кб", "label"),
      ("B20", '=B16*(1-B8/100)*(1-VLOOKUP(B7&", % сжатия",A3:B6,2,0)/100)', "value")]
S += [("A21", "вес записи в субд расжатый, кб", "label"), ("B21", "=B16*B18/100", "value")]
S += [("A22", "вес записи в субд основной, кб", "label"), ("B22", "=B21*(1-B11/100)", "value")]
S += [("A23", "вес записи в субд с учетом избыточности, кб", "label"), ("B23", "=B22*B17", "value")]
S += [("A25", "итоговые сводные ячейки - " + "источник", "sum")]
S += [("B26", "час", "sum"), ("C26", "сутки", "sum"), ("D26", "месяц", "sum"), ("E26", "год", "sum")]
S += [("A27", "записей, миллионы", "label")]
S += [("A28", "кафка записи, ТБ", "label"), ("A29", "фх записи, ТБ", "label"),
      ("A30", "кликхаус записи, ТБ", "label")]
S += [("A32", "итоговые сводные ячейки железа - " + "источник", "sum")]
for col, name in zip("BCDE", ["CPU", "RAM", "SSD", "HDD"]):
    S += [(f"{col}33", name, "sum")]
for col, name in zip("GHIJ", ["cpu на ТБ данных год", "ram ТБ на ТБ данных год",
                              "SSD ТБ на ТБ данных год", "HDD ТБ на ТБ данных год"]):
    S += [(f"{col}33", name, "sum")]
S += [("L33", "Множитель лет", "sum")]
NORMS = {34: (0.1, 0.02, 3, 0, 3), 35: (2, 0.15, 0.5, 2, 1), 36: (0.05, 0.03, 0, 1.2, 3)}
for i, (name, ycol) in enumerate([("кафка", "E28"), ("кликхаус", "E30"), ("минио", "E29")]):
    r = 34 + i
    S += [(f"A{r}", name, "label")]
    for col, norm in zip("BCDE", "GHIJ"):
        S += [(f"{col}{r}", f"=ROUNDUP({ycol}*$L{r}*{norm}{r},2)", "value")]
    for col, v in zip("GHIJL", NORMS[r]):
        S += [(f"{col}{r}", v, "value")]
S += [("A46", "бонус", "section")]
S += [("A47", "процент записи подверженный сжатию, %", "label"), ("B47", 31, "value")]
S += [("A48", "колонки под сжатие", "label"), ("B48", "lat, lon", "literal")]
S += [("D48", "кодек", "label"), ("E48", "процент сжатия кодеком, %", "label")]
S += [("D49", "кодек 1", "literal"), ("E49", 60, "value")]
S += [("D50", "кодек 2", "literal"), ("E50", 44, "value")]

for coord, value, kind in S:
    c = ws[coord]
    if value is not None:
        c.value = value
    if kind == "section":
        c.fill, c.font = FILL_SECTION, FONT_WHITE
    elif kind == "sum":
        c.fill, c.font = FILL_SUM, FONT_TEXT
    elif kind == "label":
        c.fill, c.font = FILL_LABEL, FONT_TEXT
    elif kind == "literal":
        c.font = FONT_LITERAL

R2 = ",2)"
for col, per in zip("BCDE", ("/24", "", "*30", "*365")):
    ws[f"{col}27"] = f"=ROUNDUP(B15{per}{R2}"
    ws[f"{col}28"] = f"=ROUNDUP(B15*{MN}*B16*{KB}*(1-B10/100){per}/{TB}{R2}"
    ws[f"{col}29"] = f"=ROUNDUP(B15*{MN}*B20*{KB}{per}/{TB}{R2}"
    ws[f"{col}30"] = f"=ROUNDUP(B15*{MN}*B23*{KB}*B9{per}/{TB}{R2}"

for col in range(1, 13):
    L = get_column_letter(col)
    width = max((len(str(c.value)) for row in ws.iter_rows() for c in row
                 if c.column == col and c.value is not None), default=8)
    ws.column_dimensions[L].width = min(width * 1.15 + 3, 80)

wb.save(OUT)
print(f"OK: skeleton written to {OUT}; fill values and run check.py before delivery")
