#!/usr/bin/env python3
"""Unwind final cells of the calculator: print the full computation tree - label, formula,
evaluated value - so a human can review the logic and the math behind every итоговое число.
Usage: unwind.py WORKBOOK.xlsx [CELL ...]
No cells given: unwinds every formula cell that no other formula references (the true finals)."""
import math
import re
import sys

import openpyxl

REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
RANGE = re.compile(r"(\$?[A-Z]{1,3}\$?\d+):(\$?[A-Z]{1,3}\$?\d+)")


def col_num(col):
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch) - 64)
    return n


def refs_from_formula(f):
    out = set()
    for m in RANGE.finditer(f):
        c1 = re.match(r"([A-Z]+)(\d+)", m.group(1).replace("$", ""))
        c2 = re.match(r"([A-Z]+)(\d+)", m.group(2).replace("$", ""))
        for col in range(col_num(c1.group(1)), col_num(c2.group(1)) + 1):
            for row in range(int(c1.group(2)), int(c2.group(2)) + 1):
                out.add(f"{col}{row}")
    for m in REF.finditer(f):
        out.add(f"{m.group(1)}{m.group(2)}")
    return out


w = openpyxl.load_workbook(sys.argv[1]).active
memo = {}
ru = lambda x: math.ceil(x * 100) / 100
AMP, PLUS = chr(38), chr(43)


def cell_value(ref):
    ref = ref.replace("$", "")
    if ref in memo:
        return memo[ref]
    v = w[ref].value
    if v is None:
        raise AssertionError(f"dangling reference {ref}")
    if isinstance(v, str) and v.startswith("="):
        v = evaluate(v)
    memo[ref] = v
    return v


def evaluate(formula):
    s = formula[1:]

    def vl(m):
        key_expr, rng = m.group(1), m.group(2)
        key = eval(re.sub(r'(?<![A-Z0-9$:"])(\$?[A-Z]{1,3}\d+)',
                          lambda mm: repr(w[mm.group(1).replace("$", "")].value),
                          key_expr.replace(AMP, PLUS)))
        c1, c2 = rng.split(":")
        col, r1 = re.match(r"([A-Z]+)(\d+)", c1).groups()
        r2 = int(re.match(r"[A-Z]+(\d+)", c2).group(1))
        col2 = chr(ord(col) + 1)
        for r in range(int(r1), r2 + 1):
            if str(w[f"{col}{r}"].value) == str(key):
                return str(w[f"{col2}{r}"].value)
        raise AssertionError(f"#N/A: key {key!r} not found in {rng}")

    s = re.sub(r"VLOOKUP\((.+),([A-Z]\d+:[A-Z]\d+),2,0\)", vl, s)
    s = re.sub(r"(?<![A-Z0-9:])(\$?[A-Z]{1,3}\d+)",
               lambda m: repr(cell_value(m.group(1))), s)
    s = resolve_ifs(s)
    s = re.sub(r"ROUNDUP\((.+),\s*2\)", lambda m: f"__ru({m.group(1)})", s)
    s = re.sub(r"POWER\((.+),(.+)\)", lambda m: f"({m.group(1)})**({m.group(2)})", s)
    return eval(s, {"__ru": ru})


def resolve_ifs(s):
    while "IF(" in s:
        i = s.index("IF(")
        args, cur, depth, j = [], [], 1, i + 3
        while depth:
            ch = s[j]
            if ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth == 0:
                    break
            elif ch == "," and depth == 1:
                args.append("".join(cur)); cur = []; j += 1; continue
            cur.append(ch); j += 1
        args.append("".join(cur))
        cond, a, b = args[0], args[1], args[2]
        cond = cond.replace("<>", "!=")
        cond = re.sub(r"(?<![<>=!])=(?!=)", "==", cond)
        repl = a if eval(cond) else b
        s = s[:i] + resolve_ifs(repl) + s[j + 1:]
    return s


def label_of(coord):
    m = REF.match(coord)
    row, col = int(m.group(2)), col_num(m.group(1))
    lab = ""
    v = w.cell(row, 1).value
    if isinstance(v, str) and v and not v.startswith("="):
        lab = v
    if col > 2:
        for r in range(row - 1, max(1, row - 12), -1):
            h = w.cell(r, col).value
            if isinstance(h, str) and h and not h.startswith("="):
                lab = f"{lab} / {h}" if lab else h
                break
    return lab


def tree(coord, depth=0, seen=None):
    seen = seen or set()
    cell = w[coord]
    pad = "  " * depth
    if coord in seen:
        print(f"{pad}{coord} [{label_of(coord)}] = (cycle)")
        return
    if isinstance(cell.value, str) and cell.value.startswith("="):
        try:
            val = cell_value(coord)
        except Exception as e:
            val = f"!! {e}"
        print(f"{pad}{coord} [{label_of(coord)}] = {val}   {cell.value}")
        seen = seen | {coord}
        for ref in sorted(refs_from_formula(cell.value)):
            tree(ref, depth + 1, seen)
    else:
        print(f"{pad}{coord} [{label_of(coord)}] = {cell.value!r}")


formulas = {c.coordinate: c.value for row in w.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")}
targets = sys.argv[2:]
if not targets:
    referenced = set()
    for f in formulas.values():
        referenced |= refs_from_formula(f)
    targets = [c for c in formulas if c not in referenced]
for coord in targets:
    print("=" * 60)
    tree(coord.replace("$", ""))
